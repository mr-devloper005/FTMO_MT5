"""The fixed-universe research report: quantitative leaderboard + company evidence, side by side.

Everything numeric and every verdict comes from the store (ranking.leaderboard,
the deterministic policy). Coverage comes from company.coverage (deterministic).
The agent contributes only the company-context text it recorded with
record_company_context, and the RAG Q&A comes from company_research_log. Company
evidence never changes a rank or a verdict.
"""
from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any, Dict, List, Optional

from research_agent import company, ranking
from research_agent.config import Settings
from research_agent.export import _FILL, _WIDTHS, _cell
from research_agent.store import Store

COLUMNS = ["Rank", "Stock", "Best Strategy", "Best Timeframe", "Strategy Return %", "Buy & Hold %", "Beat by (pts)",
           "Robustness", "Deterministic Verdict", "Verdict Reason", "Company Context", "RAG Evidence Coverage",
           "Reason"]


def _fmt(v: Any, pct: bool = True) -> str:
    if v is None:
        return "n/a"
    return f"{v:+.1f}%" if pct else f"{v:+.1f}"


def leaderboard(settings: Settings, run_id: str, symbols: List[str]) -> List[Dict[str, Any]]:
    best = {r["Stock"]: r for r in ranking.leaderboard(settings, limit=10_000, best_per_stock=True, symbols=symbols)}
    ctx = company.contexts(settings, run_id)
    rows = []
    for sym in symbols:
        r, c = best.get(sym), ctx.get(sym, {})
        cov = company.coverage(sym)
        rows.append({"Stock": sym, "_lb": r, "_cov": cov,
                     "Best Strategy": r["Strategy"] if r else "no result",
                     "Best Timeframe": r["Timeframe"] if r else "n/a",
                     "Strategy Return %": r["Return %"] if r else None,
                     "Buy & Hold %": r["Buy & Hold %"] if r else None,
                     "Beat by (pts)": r["Outperformance (pts)"] if r else None,
                     "Robustness": r["Robustness"] if r else "n/a",
                     "Deterministic Verdict": r["Verdict"] if r else "No result",
                     "Verdict Reason": r["Reason"] if r else "No backtest result for this stock.",
                     "Company Context": c.get("company_context") or "(not recorded)",
                     "RAG Evidence Coverage": company.coverage_line(cov),
                     "Reason": c.get("final_reason") or "(not recorded)",
                     "_order": r["Rank"] if r else 10_000})
    rows.sort(key=lambda x: x["_order"])                 # deterministic order, untouched by company context
    for i, x in enumerate(rows, 1):
        x["Rank"] = i
    return rows


def _csv(rows: List[Dict[str, Any]], path: Path) -> Path:
    with path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=COLUMNS, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow({k: _cell(r.get(k)) for k in COLUMNS})
    return path


def _xlsx(rows: List[Dict[str, Any]], qa: List[Dict[str, Any]], path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    tabs = {"Leaderboard": (COLUMNS, [[r.get(k) for k in COLUMNS] for r in rows]),
            "RAG Q&A": (["Stock", "Question", "Answer", "Not found", "Grounded", "Sources", "Cost $"],
                        [[q["symbol"], q["question"], q["answer"], bool(q["not_found"]), bool(q["grounded"]),
                          "\n".join(q["citations"] or []), q["cost_usd"]] for q in qa])}
    wb.remove(wb.active)
    for name, (cols, data) in tabs.items():
        ws = wb.create_sheet(name)
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F2A44")
        for row in data:
            ws.append([_cell(v) for v in row])
        for i, k in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(i)].width = {"Company Context": 60, "Reason": 60, "Answer": 90,
                                                                "Question": 50, "Sources": 70,
                                                                "Verdict Reason": 50}.get(k, _WIDTHS.get(k, 14))
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        if "Deterministic Verdict" in cols:
            vi = cols.index("Deterministic Verdict") + 1
            for r in range(2, ws.max_row + 1):
                fill = _FILL.get(ws.cell(row=r, column=vi).value)
                if fill:
                    ws.cell(row=r, column=vi).fill = PatternFill("solid", fgColor=fill)
        ws.freeze_panes = "C2"
    wb.save(path)
    return path


def _md_table(rows: List[Dict[str, Any]]) -> str:
    head = ["Rank", "Stock", "Best Strategy", "Best Timeframe", "Strategy Return", "Buy & Hold", "Beat by",
            "Robustness", "Deterministic Verdict", "Company Context", "RAG Evidence Coverage", "Reason"]
    out = ["| " + " | ".join(head) + " |", "|" + "---|" * len(head)]
    for r in rows:
        cells = [r["Rank"], r["Stock"], r["Best Strategy"], r["Best Timeframe"], _fmt(r["Strategy Return %"]),
                 _fmt(r["Buy & Hold %"]), _fmt(r["Beat by (pts)"], pct=False) + (" pts" if r["Beat by (pts)"] is not None else ""),
                 r["Robustness"], r["Deterministic Verdict"], r["Company Context"], r["RAG Evidence Coverage"],
                 r["Reason"]]
        out.append("| " + " | ".join(str(c).replace("|", "/").replace("\n", " ") for c in cells) + " |")
    return "\n".join(out)


def _stock_section(r: Dict[str, Any], ctx: Dict[str, Any], qa: List[Dict[str, Any]]) -> str:
    lb, cov = r["_lb"], r["_cov"]
    lines = [f"### {r['Stock']}", "", "**Quantitative (deterministic):**", ""]
    if lb:
        lines += [f"- Best strategy: {lb['Strategy']} ({lb['Source']})", f"- Best timeframe: {lb['Timeframe']}",
                  f"- Return: {_fmt(lb['Return %'])}  |  Buy & Hold: {_fmt(lb['Buy & Hold %'])}  |  "
                  f"Beat by: {_fmt(lb['Outperformance (pts)'], pct=False)} pts",
                  f"- Robustness: {lb['Robustness']}  |  Trades: {lb['Trades']}  |  Max DD: {_fmt(lb['Max DD %'])}"
                  f"  |  Years: {lb['Years']}  |  Data as of: {lb['Data as of']}",
                  f"- Verdict: **{lb['Verdict']}** — {lb['Reason']}"]
    else:
        lines.append("- No backtest result.")
    lines += ["", "**Company evidence (RAG, read-only):**", "",
              f"- Financial condition: {ctx.get('financial_condition') or '(not recorded)'}",
              f"- Risks: {ctx.get('risks') or '(not recorded)'}",
              f"- Guidance / outlook: {ctx.get('guidance') or '(not recorded)'}",
              f"- Management evidence: {ctx.get('management') or '(not recorded)'}", "",
              f"**Evidence quality:** {cov['evidence_quality']} — {company.coverage_line(cov)}", ""]
    if cov["missing"]:
        lines.append("Missing evidence:")
        lines += [f"- {m}" for m in cov["missing"]]
        lines.append("")
    lines += [f"**Final reason:** {ctx.get('final_reason') or '(not recorded)'}", "",
              f"RAG questions asked for {r['Stock']}: {len(qa)} "
              f"({sum(1 for q in qa if q['not_found'])} not found, "
              f"{sum(1 for q in qa if q['grounded'])} grounded)", ""]
    return "\n".join(lines)


def _qa_section(qa: List[Dict[str, Any]]) -> str:
    out = []
    for q in qa:
        ev = q["evidence"] or []
        out += [f"#### {q['symbol']}: {q['question']}", "",
                f"- Route: structured={q['route'].get('structured') if q['route'] else '?'}, "
                f"documents={q['route'].get('documents') if q['route'] else '?'}",
                "- Retrieved evidence:" if ev else "- Retrieved evidence: none"]
        out += [f"  - [{e['n']}] {e['source']} — \"{e['excerpt'][:180]}…\"" for e in ev[:3]]
        if len(ev) > 3:
            out.append(f"  - … {len(ev) - 3} more")
        out += ["- Sources cited: " + ("; ".join(q["citations"]) if q["citations"] else "none"),
                "- Claude answer:", "", "> " + q["answer"].replace("\n", "\n> "), "",
                f"- Grounded: {'yes' if q['grounded'] else 'NO'}  |  Not found: {'yes' if q['not_found'] else 'no'}"
                f"  |  number guard unverified: {(q['guard'] or {}).get('unverified') or 'none'}"
                f"  |  database conflicts: {len((q['factcheck'] or {}).get('conflicts') or [])}", ""]
    return "\n".join(out)


def build(settings: Settings, run_id: str, symbols: List[str], run_stats: Dict[str, Any],
          rag_assessment: str = "") -> Dict[str, Any]:
    settings.ensure_dirs()
    rows = leaderboard(settings, run_id, symbols)
    ctx = company.contexts(settings, run_id)
    qa = company.research_log(settings, run_id)
    by_sym = {s: [q for q in qa if q["symbol"] == s] for s in symbols}
    store = Store(settings.db_path)
    screens = {s: store.screen(s) or {} for s in symbols}
    md = [f"# FTMO AI Research — fixed 5-stock run `{run_id}`", "",
          "Research only: no broker, order, position or execution access. Verdicts, ranks and robustness come "
          "from the deterministic FTMO policy; company evidence is shown beside them and never changes them.", "",
          f"Stocks: {', '.join(symbols)}  |  Timeframes ({len(settings.timeframes)}): "
          f"{', '.join(settings.timeframes)}", "",
          "## Final leaderboard", "", _md_table(rows), "",
          "## Run statistics", "", "```", json.dumps(run_stats, indent=1, default=str), "```", "",
          "## Market / data context (screen, daily bars)", "",
          "| Stock | Data as of | Passed screen | Reason | Price | ATR % | 1y return % | History bars |",
          "|---|---|---|---|---|---|---|---|"]
    for s in symbols:
        x = screens[s]
        md.append(f"| {s} | {x.get('as_of')} | {'yes' if x.get('passed') else 'no'} | {x.get('reason')} | "
                  f"{x.get('price')} | {x.get('atr_pct')} | {x.get('return_1y')} | {x.get('history_bars')} |")
    md += ["", "## Company research (per stock)", ""]
    md += [_stock_section(r, ctx.get(r["Stock"], {}), by_sym[r["Stock"]]) for r in rows]
    md += ["## RAG validation (every question asked this run)", "", _qa_section(qa), "",
           "## Does RAG add company-level context? (agent assessment)", "",
           rag_assessment.strip() or "(not recorded)", ""]
    base = settings.reports_dir / f"{run_id}-fixed-5-stock-research"
    paths = {"markdown": str(base.with_suffix(".md"))}
    Path(paths["markdown"]).write_text("\n".join(md), encoding="utf-8")
    exp = settings.exports_dir / f"research-{run_id}"
    paths["csv"] = str(_csv(rows, exp.with_suffix(".csv")))
    try:
        paths["xlsx"] = str(_xlsx(rows, qa, exp.with_suffix(".xlsx")))
    except ImportError:
        paths["xlsx_error"] = "openpyxl not installed"
    return {"paths": paths, "rows": [{k: v for k, v in r.items() if not k.startswith("_")} for r in rows],
            "rag_questions": len(qa), "rag_cost_usd": round(sum(q["cost_usd"] or 0 for q in qa), 4)}
